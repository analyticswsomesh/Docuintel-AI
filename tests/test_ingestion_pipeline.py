from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document as DocxDocument
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.db.base import Base
from app.db.repositories import DocumentRepository
from app.services.chunking_service import ChunkingError, TextChunkingService
from app.services.extraction_service import DocumentExtractionService, ExtractedPage, ExtractionError
from app.services.ingestion_service import DocumentIngestionError, DocumentIngestionService, DocumentStatus
from app.services.vector_store_service import TextChunk


def make_repository() -> tuple[Session, DocumentRepository]:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = Session(engine)
    return session, DocumentRepository(session)


class FakeEmbeddingService:
    def embed(self, texts):
        return [[float(index), 1.0] for index, _ in enumerate(texts)]


class FakeVectorStore:
    def __init__(self, fail_upsert: bool = False):
        self.deleted_document_ids = []
        self.upserted: list[tuple[list[TextChunk], list[list[float]]]] = []
        self.fail_upsert = fail_upsert
        self.events = []

    def delete_document(self, document_id: str) -> None:
        self.deleted_document_ids.append(document_id)
        self.events.append(("delete", document_id))

    def upsert(self, chunks, embeddings) -> None:
        if self.fail_upsert:
            raise RuntimeError("vector write failed")
        self.upserted.append((chunks, embeddings))
        self.events.append(("upsert", chunks[0].document_id))


class FakeExtractor:
    def __init__(self, pages):
        self.pages = pages

    def extract(self, file_path, *, filename, content_type):
        return self.pages


def make_service(repository, extractor, vector_store):
    return DocumentIngestionService(
        repository=repository,
        extractor=extractor,
        chunker=TextChunkingService(chunk_size_chars=20, chunk_overlap_chars=5),
        embedding_service=FakeEmbeddingService(),
        vector_store=vector_store,
        max_upload_size_bytes=10_000,
    )


def test_txt_extraction_and_chunking_preserve_source_metadata(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("Alpha beta gamma. Delta epsilon.", encoding="utf-8")

    pages = DocumentExtractionService().extract(path, filename="notes.txt", content_type="text/plain")
    chunks = TextChunkingService(chunk_size_chars=12, chunk_overlap_chars=3).create_chunks(
        pages, document_id="doc-1", source_name="notes.txt"
    )

    assert pages[0].page_number is None
    assert chunks[0].document_id == "doc-1"
    assert chunks[0].source_name == "notes.txt"
    assert chunks[0].chunk_index == 0
    assert len(chunks) > 1


def test_invalid_text_input_preserves_original_decode_error(tmp_path: Path) -> None:
    path = tmp_path / "invalid.txt"
    path.write_bytes(b"valid prefix\xff")

    with pytest.raises(ExtractionError) as error:
        DocumentExtractionService().extract(path, filename="invalid.txt", content_type="text/plain")

    assert isinstance(error.value.__cause__, UnicodeDecodeError)


def test_unsupported_file_format_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "notes.csv"
    path.write_text("not supported", encoding="utf-8")

    with pytest.raises(ExtractionError, match="Unsupported document type"):
        DocumentExtractionService().extract(path)


def test_malformed_docx_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a real docx")

    with pytest.raises(ExtractionError, match="Could not read DOCX document"):
        DocumentExtractionService().extract(path)


def test_chunk_overlap_is_explicit() -> None:
    chunks = TextChunkingService(chunk_size_chars=10, chunk_overlap_chars=3).create_chunks(
        [ExtractedPage("abcdefghijklmno")], document_id="doc-1", source_name="notes.txt"
    )

    assert chunks[0].text[-3:] == chunks[1].text[:3]
    assert [chunk.chunk_index for chunk in chunks] == list(range(len(chunks)))


def test_docx_extraction_reads_paragraph_text(tmp_path: Path) -> None:
    path = tmp_path / "brief.docx"
    document = DocxDocument()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    document.save(path)

    pages = DocumentExtractionService().extract(path, filename="brief.docx", content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    assert len(pages) == 1
    assert pages[0].page_number is None
    assert "First paragraph." in pages[0].text
    assert "Second paragraph." in pages[0].text


def test_pdf_extraction_preserves_page_numbers(tmp_path: Path) -> None:
    fitz = pytest.importorskip("fitz", reason="PyMuPDF is required for PDF fixture coverage")
    path = tmp_path / "report.pdf"
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Page one evidence")
    pdf.new_page().insert_text((72, 72), "Page two evidence")
    pdf.save(path)
    pdf.close()

    pages = DocumentExtractionService().extract(path, filename="report.pdf", content_type="application/pdf")

    assert [page.page_number for page in pages] == [1, 2]
    assert "Page one evidence" in pages[0].text


def test_empty_text_is_rejected() -> None:
    with pytest.raises(ChunkingError):
        TextChunkingService().create_chunks([ExtractedPage("  ")], document_id="doc-1", source_name="empty.txt")


def test_ingestion_transitions_to_indexed_and_replaces_vectors(tmp_path: Path) -> None:
    session, repository = make_repository()
    try:
        document = repository.create(filename="notes.txt", content_type="text/plain", size_bytes=10, checksum="c" * 64)
        path = tmp_path / "stored-file.txt"
        path.write_text("source", encoding="utf-8")
        vector_store = FakeVectorStore()
        service = make_service(repository, FakeExtractor([ExtractedPage("A useful passage", 3)]), vector_store)

        result = service.ingest(document.id, path)

        assert result.status == DocumentStatus.INDEXED.value
        assert result.error_message is None
        assert vector_store.deleted_document_ids == [str(document.id)]
        assert vector_store.upserted[0][0][0].page_number == 3
        assert vector_store.upserted[0][0][0].document_id == str(document.id)
    finally:
        session.close()


def test_ingestion_failure_clears_vectors_and_marks_document_failed(tmp_path: Path) -> None:
    session, repository = make_repository()
    try:
        document = repository.create(filename="notes.txt", content_type="text/plain", size_bytes=10, checksum="d" * 64)
        path = tmp_path / "stored-file.txt"
        path.write_text("source", encoding="utf-8")
        vector_store = FakeVectorStore(fail_upsert=True)
        service = make_service(repository, FakeExtractor([ExtractedPage("A useful passage")]), vector_store)

        with pytest.raises(DocumentIngestionError):
            service.ingest(document.id, path)

        failed = repository.get(document.id)
        assert failed.status == DocumentStatus.FAILED.value
        assert failed.error_message == "vector write failed"
        assert vector_store.deleted_document_ids == [str(document.id), str(document.id)]
    finally:
        session.close()


def test_reprocessing_clears_existing_chunks_before_writing_replacement(tmp_path: Path) -> None:
    session, repository = make_repository()
    try:
        document = repository.create(filename="notes.txt", content_type="text/plain", size_bytes=10, checksum="e" * 64)
        path = tmp_path / "stored-file.txt"
        path.write_text("source", encoding="utf-8")
        vector_store = FakeVectorStore()
        service = make_service(repository, FakeExtractor([ExtractedPage("A useful passage")]), vector_store)

        service.ingest(document.id, path)
        service.ingest(document.id, path)

        assert vector_store.events == [
            ("delete", str(document.id)),
            ("upsert", str(document.id)),
            ("delete", str(document.id)),
            ("upsert", str(document.id)),
        ]
    finally:
        session.close()


def test_ingestion_rejects_unknown_document_id(tmp_path: Path) -> None:
    session, repository = make_repository()
    try:
        service = make_service(repository, FakeExtractor([]), FakeVectorStore())
        with pytest.raises(DocumentIngestionError, match="not found"):
            service.ingest(uuid4(), tmp_path / "missing.txt")
    finally:
        session.close()
